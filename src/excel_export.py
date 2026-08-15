"""Excel export for validated TestScope AI test packs."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from src.schemas import TestPack


HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)
HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)


def style_worksheet(
    worksheet: Worksheet,
    column_widths: dict[str, int],
) -> None:
    """Apply consistent formatting to an exported worksheet."""

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width


def add_summary_sheet(
    workbook: Workbook,
    test_pack: TestPack,
) -> None:
    """Add high-level requirement and TestPack information."""

    worksheet = workbook.active
    worksheet.title = "Summary"

    worksheet.append(["Metric", "Value"])
    worksheet.append(
        ["Requirement ID", test_pack.requirement.requirement_id]
    )
    worksheet.append(["Title", test_pack.requirement.title])
    worksheet.append(["Domain", test_pack.requirement.domain or ""])
    worksheet.append(["Feature", test_pack.requirement.feature or ""])
    worksheet.append(["User story", test_pack.requirement.user_story])
    worksheet.append(
        [
            "Acceptance criteria",
            len(test_pack.requirement.acceptance_criteria),
        ]
    )
    worksheet.append(
        ["Generated UAT tests", len(test_pack.test_cases)]
    )
    worksheet.append(
        [
            "Coverage percentage",
            test_pack.coverage_summary.coverage_percentage / 100,
        ]
    )
    worksheet.append(
        [
            "Requirement summary",
            test_pack.analysis.requirement_summary,
        ]
    )

    worksheet["B9"].number_format = "0.0%"

    style_worksheet(
        worksheet,
        {
            "A": 28,
            "B": 100,
        },
    )


def add_test_cases_sheet(
    workbook: Workbook,
    test_pack: TestPack,
) -> None:
    """Add one reviewable row per generated UAT test."""

    worksheet = workbook.create_sheet("UAT Test Cases")

    worksheet.append(
        [
            "Test ID",
            "Title",
            "Requirement ID",
            "Acceptance Criteria",
            "Test Type",
            "Priority",
            "Risk Level",
            "Status",
            "Objective",
            "Preconditions",
            "Test Data",
            "Test Steps",
            "Expected Result",
            "Risk IDs",
            "Ambiguity IDs",
            "Assumption IDs",
        ]
    )

    for test_case in test_pack.test_cases:
        steps = "\n".join(
            f"{step.step_number}. {step.action}"
            for step in test_case.steps
        )

        worksheet.append(
            [
                test_case.test_id,
                test_case.title,
                test_case.requirement_id,
                ", ".join(test_case.acceptance_criteria_ids),
                str(test_case.test_type),
                str(test_case.priority),
                str(test_case.risk_level),
                str(test_case.status),
                test_case.objective,
                "\n".join(test_case.preconditions),
                "\n".join(test_case.test_data),
                steps,
                test_case.expected_result,
                ", ".join(test_case.risk_ids),
                ", ".join(test_case.ambiguity_ids),
                ", ".join(test_case.assumption_ids),
            ]
        )

    style_worksheet(
        worksheet,
        {
            "A": 12,
            "B": 38,
            "C": 16,
            "D": 22,
            "E": 16,
            "F": 12,
            "G": 12,
            "H": 22,
            "I": 42,
            "J": 42,
            "K": 36,
            "L": 65,
            "M": 55,
            "N": 18,
            "O": 18,
            "P": 18,
        },
    )


def add_analysis_sheet(
    workbook: Workbook,
    test_pack: TestPack,
) -> None:
    """Add rules, ambiguities, assumptions and risks."""

    worksheet = workbook.create_sheet("Requirement Analysis")

    worksheet.append(
        [
            "Finding Type",
            "Finding ID",
            "Classification",
            "Description",
            "Related Criteria",
            "Question / Rationale",
        ]
    )

    for rule in test_pack.analysis.business_rules:
        worksheet.append(
            [
                "Business Rule",
                rule.rule_id,
                "",
                rule.description,
                ", ".join(rule.source_criteria),
                "",
            ]
        )

    for ambiguity in test_pack.analysis.ambiguities:
        worksheet.append(
            [
                "Ambiguity",
                ambiguity.ambiguity_id,
                ambiguity.impact,
                ambiguity.description,
                ", ".join(ambiguity.related_criteria),
                ambiguity.clarification_question,
            ]
        )

    for assumption in test_pack.analysis.assumptions:
        worksheet.append(
            [
                "Assumption",
                assumption.assumption_id,
                str(assumption.status),
                assumption.description,
                ", ".join(assumption.related_criteria),
                (
                    "Human approval required"
                    if assumption.approval_required
                    else "Approval not required"
                ),
            ]
        )

    for risk in test_pack.analysis.business_risks:
        worksheet.append(
            [
                "Business Risk",
                risk.risk_id,
                str(risk.severity),
                risk.description,
                ", ".join(risk.related_criteria),
                risk.rationale,
            ]
        )

    style_worksheet(
        worksheet,
        {
            "A": 20,
            "B": 16,
            "C": 25,
            "D": 60,
            "E": 22,
            "F": 65,
        },
    )


def add_coverage_sheet(
    workbook: Workbook,
    test_pack: TestPack,
) -> None:
    """Add deterministic coverage results."""

    coverage = test_pack.coverage_summary
    worksheet = workbook.create_sheet("Coverage")

    worksheet.append(["Metric", "Value"])
    worksheet.append(["Total criteria", coverage.total_criteria])
    worksheet.append(
        ["Covered criteria", ", ".join(coverage.covered_criteria)]
    )
    worksheet.append(
        ["Uncovered criteria", ", ".join(coverage.uncovered_criteria)]
    )
    worksheet.append(
        ["Coverage percentage", coverage.coverage_percentage / 100]
    )
    worksheet.append(
        ["High-risk tests", ", ".join(coverage.high_risk_tests)]
    )

    worksheet.append([])
    worksheet.append(["Test Type", "Number of Tests"])

    for test_type, count in coverage.tests_by_type.items():
        worksheet.append([str(test_type), count])

    worksheet["B5"].number_format = "0.0%"

    style_worksheet(
        worksheet,
        {
            "A": 30,
            "B": 80,
        },
    )


def create_test_pack_excel(
    test_pack: TestPack,
) -> bytes:
    """Return a formatted TestPack workbook as XLSX bytes."""

    workbook = Workbook()

    add_summary_sheet(workbook, test_pack)
    add_test_cases_sheet(workbook, test_pack)
    add_analysis_sheet(workbook, test_pack)
    add_coverage_sheet(workbook, test_pack)

    output = BytesIO()
    workbook.save(output)

    return output.getvalue()
