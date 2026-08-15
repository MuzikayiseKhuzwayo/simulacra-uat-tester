"""Tests for the TestPack Excel exporter."""

from io import BytesIO

from openpyxl import load_workbook

from src.excel_export import create_test_pack_excel
from tests.test_schemas import valid_test_pack


def load_exported_workbook():
    """Generate the Excel file and load it for verification."""

    excel_bytes = create_test_pack_excel(valid_test_pack())
    return load_workbook(BytesIO(excel_bytes))


def test_excel_export_contains_expected_sheets() -> None:
    """The workbook must contain all required business sheets."""

    workbook = load_exported_workbook()

    assert workbook.sheetnames == [
        "Summary",
        "UAT Test Cases",
        "Requirement Analysis",
        "Coverage",
    ]


def test_excel_export_contains_generated_test_case() -> None:
    """The UAT worksheet must contain the generated test case."""

    workbook = load_exported_workbook()
    worksheet = workbook["UAT Test Cases"]

    assert worksheet["A1"].value == "Test ID"
    assert worksheet["A2"].value == "UAT-001"
    assert worksheet["B2"].value == "Complete a transfer below the daily limit"


def test_excel_export_contains_numeric_coverage() -> None:
    """Coverage must be stored as a numeric percentage."""

    workbook = load_exported_workbook()
    worksheet = workbook["Summary"]

    assert worksheet["A9"].value == "Coverage percentage"
    assert worksheet["B9"].value == 1
    assert worksheet["B9"].number_format == "0.0%"